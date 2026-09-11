"""JSON → documentos Word con el formato de Mentes Brillantes."""

import io
import json
import zipfile
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
GOLD = RGBColor(0xC9, 0xA9, 0x61)


# --------------------------------------------------------------------- utils

def _doc(margen=0.9) -> Document:
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Inches(margen)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    return d


def _h(d, texto, nivel=1):
    p = d.add_heading(str(texto), level=nivel)
    for r in p.runs:
        r.font.color.rgb = NAVY
    return p


def _p(d, texto, size=11):
    for bloque in str(texto or "").split("\n\n"):
        bloque = bloque.strip()
        if not bloque:
            continue
        bold = bloque.startswith("**")
        bloque = bloque.replace("**", "")
        par = d.add_paragraph()
        r = par.add_run(bloque)
        r.bold = bold
        r.font.size = Pt(size)
        par.paragraph_format.space_after = Pt(8)


def _kv(d, clave, valor, italic=False):
    par = d.add_paragraph()
    r = par.add_run(f"{clave}: ")
    r.bold = True
    r2 = par.add_run(str(valor))
    r2.italic = italic


def _bul(d, items):
    for i in items or []:
        d.add_paragraph(str(i), style="List Bullet")


def _celda(c, texto, bold=False, size=9):
    c.text = ""
    p = c.paragraphs[0]
    r = p.add_run(str(texto))
    r.bold = bold
    r.font.size = Pt(size)


def _guardar(d) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------- módulo

def modulo_docx(m: dict, curso: str, codigo: str) -> bytes:
    d = _doc()
    d.add_heading(curso, 0)
    p = d.add_paragraph()
    r = p.add_run(f"{codigo} · Club de Mentes Brillantes")
    r.italic = True
    d.add_paragraph(f"MÓDULO {m.get('numero')}. {str(m.get('nombre', '')).upper()}")
    _kv(d, "Duración", f"{m.get('duracion_horas')} horas")
    _kv(d, "Objetivo del módulo", m.get("objetivo", ""))
    d.add_page_break()

    _h(d, "Introducción al módulo")
    _p(d, m.get("introduccion"))

    v = m.get("video") or {}
    _h(d, "Video del módulo")
    _kv(d, "Título sugerido", v.get("titulo_sugerido", ""))
    _kv(d, "Duración", f"{v.get('duracion_sugerida_min', '')} minutos")
    _p(d, v.get("guion_resumen"))
    d.add_paragraph("[ESPACIO RESERVADO PARA EL VIDEO]")

    f = m.get("foro") or {}
    _h(d, "Foro de discusión")
    _p(d, f.get("pregunta"))
    _kv(d, "Instrucción", f.get("instruccion", ""))

    for t in m.get("temas") or []:
        d.add_page_break()
        _h(d, f"Tema {t.get('numero')}. {t.get('nombre')}")
        _kv(d, "Objetivo del tema", t.get("objetivo", ""))

        _h(d, "Introducción", 2)
        _p(d, t.get("introduccion"))

        if t.get("preguntas_previas"):
            _h(d, "Antes de leer", 2)
            d.add_paragraph(
                "Intenta responder estas preguntas antes de continuar. No se califican: "
                "sirven para que el contenido que sigue se fije mejor.")
            _bul(d, t["preguntas_previas"])

        cb = t.get("conceptos_basicos") or {}
        _h(d, "Conceptos básicos", 2)
        _p(d, cb.get("nucleo"))

        if cb.get("conceptos_clave"):
            _h(d, "Conceptos clave", 3)
            tab = d.add_table(rows=1, cols=2)
            tab.style = "Table Grid"
            _celda(tab.rows[0].cells[0], "Término", True)
            _celda(tab.rows[0].cells[1], "Definición", True)
            for c in cb["conceptos_clave"]:
                row = tab.add_row()
                _celda(row.cells[0], c.get("termino", ""), True)
                _celda(row.cells[1], c.get("definicion", ""))
            d.add_paragraph()

        dia = cb.get("diagrama") or {}
        if dia.get("descripcion"):
            _h(d, "Diagrama", 3)
            _kv(d, "Tipo", dia.get("tipo", "esquema"))
            _p(d, dia["descripcion"])
            d.add_paragraph("[ESPACIO RESERVADO PARA EL DIAGRAMA]")

        if cb.get("modelos_o_herramientas"):
            _h(d, "Modelos y herramientas", 3)
            _bul(d, cb["modelos_o_herramientas"])

        if cb.get("ejemplos"):
            _h(d, "Ejemplos", 3)
            _bul(d, cb["ejemplos"])

        if cb.get("profundiza"):
            _h(d, "Profundiza (opcional)", 3)
            par = d.add_paragraph()
            r = par.add_run(
                "Este bloque amplía lo anterior. En Moodle va colapsado: "
                "no es necesario para la práctica ni para la evaluación.")
            r.italic = True
            r.font.size = Pt(9)
            _p(d, cb["profundiza"], size=10)

        co = t.get("conecta") or {}
        _h(d, "Conecta", 2)
        _p(d, co.get("texto"))
        if co.get("wiifm"):
            _kv(d, "¿Qué hay para mí?", co["wiifm"])
        if co.get("frase_memorable"):
            _kv(d, "Frase memorable", co["frase_memorable"], italic=True)

        pr = t.get("practica") or {}
        _h(d, "Practica", 2)
        par = d.add_paragraph()
        r = par.add_run(str(pr.get("nombre", "")))
        r.bold = True
        _kv(d, "Objetivo", pr.get("objetivo", ""))
        _h(d, "Instrucciones", 3)
        _p(d, pr.get("instrucciones"))
        _h(d, "Preguntas detonantes", 3)
        _bul(d, pr.get("preguntas_detonantes"))
        _kv(d, "Entregable", pr.get("entregable", ""))
        _kv(d, "Duración estimada", f"{pr.get('duracion_min', '')} minutos")
        if pr.get("herramienta_sugerida"):
            _kv(d, "Herramienta sugerida", pr["herramienta_sugerida"])
        if pr.get("cuenta_para_evaluacion_final"):
            _kv(d, "Nota", "Esta práctica cuenta para la evaluación final del módulo.")

        ad = t.get("adopta") or {}
        _h(d, "Adopta", 2)
        _p(d, ad.get("texto"))
        if ad.get("preguntas_reflexion"):
            _h(d, "Preguntas de reflexión", 3)
            _bul(d, ad["preguntas_reflexion"])
        if ad.get("takeaway"):
            _kv(d, "Takeaway", ad["takeaway"], italic=True)

    d.add_page_break()
    c = m.get("cierre") or {}
    _h(d, "Cierre del módulo")
    _p(d, c.get("texto"))
    if c.get("incorporacion_vida_real"):
        _kv(d, "Cómo incorporarlo", c["incorporacion_vida_real"])
    if c.get("takeaway_final"):
        _kv(d, "Takeaway final", c["takeaway_final"], italic=True)
    if c.get("cierre_wow"):
        _kv(d, "Cierre WOW", c["cierre_wow"], italic=True)

    ev = m.get("evaluacion_placeholder") or {}
    _h(d, "Evaluación del módulo")
    d.add_paragraph(str(ev.get("titulo", "")))
    d.add_paragraph(f"{ev.get('num_preguntas', 5)} preguntas de opción múltiple.")
    d.add_paragraph("[ESPACIO RESERVADO — ver documento de evaluaciones]")
    return _guardar(d)


# ------------------------------------------------------------------ syllabus

def syllabus_docx(s: dict) -> bytes:
    d = _doc(0.7)
    d.styles["Normal"].font.size = Pt(9.5)
    i = s.get("identificacion") or {}

    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    _celda(t.rows[0].cells[0], "Club de Mentes Brillantes", True)
    _celda(t.rows[0].cells[1], "Syllabus", True)
    _celda(t.rows[0].cells[2], "Formato Syllabus Cursos CMB 2025-26_1.2", True)
    d.add_paragraph()

    t = d.add_table(rows=4, cols=4)
    t.style = "Table Grid"
    filas = [("Curso", i.get("curso", ""), "Código", i.get("codigo", "")),
             ("Área", i.get("area", ""), "Tipo de Formación", i.get("tipo_formacion", "")),
             ("Duración", f"{i.get('duracion_horas', '')} horas", "Nivel", i.get("nivel", ""))]
    for r, (k1, v1, k2, v2) in enumerate(filas):
        _celda(t.rows[r].cells[0], k1, True)
        _celda(t.rows[r].cells[1], v1)
        _celda(t.rows[r].cells[2], k2, True)
        _celda(t.rows[r].cells[3], v2)
    _celda(t.rows[3].cells[0], "Frase del curso", True)
    t.rows[3].cells[1].merge(t.rows[3].cells[3])
    _celda(t.rows[3].cells[1], f"“{i.get('frase_curso', '')}”")
    d.add_paragraph()

    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"

    def fila(k, v):
        row = t.add_row()
        _celda(row.cells[0], k, True)
        _celda(row.cells[1], v)
        return row.cells[1]

    fila("Objetivo", s.get("objetivo", ""))
    fila("Eje estratégico", s.get("eje_estrategico", ""))

    pi = s.get("perfil_ingreso") or {}
    c = fila("Perfil de Ingreso", "")
    for a in pi.get("audiencias", []):
        c.add_paragraph(f"• {a}")
    if pi.get("conocimientos_previos"):
        c.add_paragraph(str(pi["conocimientos_previos"]))

    comp = s.get("competencias") or {}
    c = fila("Competencias y habilidades a desarrollar", "")
    for etiqueta, clave in (("Competencias técnicas", "tecnicas"), ("Habilidades blandas", "blandas")):
        p = c.add_paragraph()
        r = p.add_run(etiqueta)
        r.bold = True
        for x in comp.get(clave, []):
            c.add_paragraph(f"• {x}")

    ben = s.get("beneficios") or {}
    c = fila("Beneficios esperados", "")
    for etiqueta, clave in (("El participante desarrollará:", "participante"),
                            ("La empresa se beneficiará con:", "empresa"),
                            ("El usuario final se beneficiará con:", "usuario_final")):
        p = c.add_paragraph()
        r = p.add_run(etiqueta)
        r.bold = True
        for x in ben.get(clave, []):
            c.add_paragraph(f"• {x}")

    fila("Descripción del curso", s.get("descripcion", ""))

    b = s.get("badge") or {}
    c = fila("Badge", "")
    for e in b.get("emisores", []):
        c.add_paragraph(f"• {e}")
    p = c.add_paragraph()
    r = p.add_run("Ejes de la insignia:")
    r.bold = True
    for e in b.get("ejes", []):
        c.add_paragraph(f"• {e}")
    if b.get("descripcion_visual"):
        c.add_paragraph(str(b["descripcion_visual"]))

    c = fila("Certificaciones relacionadas adicionales", "")
    for x in s.get("certificaciones_relacionadas", []):
        c.add_paragraph(f"• {x}")
    d.add_paragraph()

    _h(d, "MÓDULOS DEL PROGRAMA")
    t = d.add_table(rows=0, cols=2)
    t.style = "Table Grid"
    for m in s.get("modulos_resumen", []):
        row = t.add_row()
        _celda(row.cells[0], f"{m.get('numero')}. {m.get('nombre')}")
        _celda(row.cells[1], str(m.get("horas", "")))
    row = t.add_row()
    _celda(row.cells[0], "")
    _celda(row.cells[1], f"HRS TOTALES {s.get('horas_totales', '')}", True)
    d.add_paragraph()

    for titulo, clave in (("HABILIDADES ASOCIADAS", "habilidades_asociadas"),):
        t = d.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        _celda(t.rows[0].cells[0], titulo, True)
        _celda(t.rows[0].cells[1], "")
        for n, x in enumerate(s.get(clave, []), 1):
            row = t.add_row()
            _celda(row.cells[0], f"{n}.")
            _celda(row.cells[1], x)
        d.add_paragraph()

    if s.get("vinculacion_otros_cursos"):
        t = d.add_table(rows=1, cols=2)
        t.style = "Table Grid"
        _celda(t.rows[0].cells[0], "VINCULACIÓN CON OTROS CURSOS", True)
        _celda(t.rows[0].cells[1], "")
        for n, v in enumerate(s["vinculacion_otros_cursos"], 1):
            row = t.add_row()
            _celda(row.cells[0], f"{n}.")
            _celda(row.cells[1], f"{v.get('curso')} ({v.get('relacion')})")
    d.add_page_break()

    for m in s.get("modulos", []):
        t = d.add_table(rows=2, cols=5)
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        _celda(hdr[0], f"MÓDULO {m.get('numero')}\n{m.get('nombre')}\n\n"
                       f"Objetivo: {m.get('objetivo', '')}\n\n"
                       f"Producto: {m.get('producto_del_modulo', '')}", True, 8.5)
        for idx, txt in enumerate(["APRENDIZAJES ESPERADOS", "CRITERIOS EVALUACIÓN",
                                   "CONTENIDOS", "HRS"], 1):
            _celda(hdr[idx], txt, True, 8.5)

        body = t.rows[1].cells
        for cell, items in ((body[0], m.get("aprendizajes_esperados")),
                            (body[1], m.get("criterios_evaluacion"))):
            cell.text = ""
            for k, x in enumerate(items or []):
                p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                r = p.add_run(f"• {x}")
                r.font.size = Pt(8)

        body[2].text = ""
        primero = True
        for cont in m.get("contenidos", []):
            p = body[2].paragraphs[0] if primero else body[2].add_paragraph()
            primero = False
            r = p.add_run(str(cont.get("tema", "")))
            r.bold = True
            r.font.size = Pt(8.5)
            for pt in cont.get("puntos", []):
                pp = body[2].add_paragraph()
                rr = pp.add_run(f"• {pt}")
                rr.font.size = Pt(8)
        _celda(body[4], str(m.get("horas", "")))

        row = t.add_row()
        _celda(row.cells[0], "Bibliografía / Recursos didácticos", True, 8.5)
        merged = row.cells[1].merge(row.cells[4])
        merged.text = ""
        for k, x in enumerate(m.get("bibliografia_recursos") or []):
            p = merged.paragraphs[0] if k == 0 else merged.add_paragraph()
            r = p.add_run(f"• {x}")
            r.font.size = Pt(8)

        row = t.add_row()
        _celda(row.cells[0], "Carpeta de evidencias / productos", True, 8.5)
        merged = row.cells[1].merge(row.cells[4])
        merged.text = ""
        for k, x in enumerate(m.get("evidencias") or []):
            p = merged.paragraphs[0] if k == 0 else merged.add_paragraph()
            r = p.add_run(f"• {x}")
            r.font.size = Pt(8)
        d.add_paragraph()

    for titulo, clave, subclaves in (
        ("ESTRATEGIAS METODOLÓGICAS", "estrategias_metodologicas",
         ["aprendizaje_activo", "evaluacion_continua", "personalizacion", "multimedia"]),
        ("ESTRATEGIAS EVALUATIVAS", "estrategias_evaluativas",
         ["diagnostica", "formativa", "sumativa", "retroalimentacion_participantes"]),
    ):
        _h(d, titulo)
        blk = s.get(clave) or {}
        t = d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for sc in subclaves:
            row = t.add_row()
            _celda(row.cells[0], sc.replace("_", " ").title(), True)
            _celda(row.cells[1], blk.get(sc, ""))
        d.add_paragraph()

    req = (s.get("estrategias_evaluativas") or {}).get("requisitos_acreditacion") or []
    if req:
        _kv(d, "Requisitos de acreditación", " · ".join(req))
        d.add_paragraph()

    _h(d, "RECURSOS EDUCATIVOS")
    rec = s.get("recursos_educativos") or {}
    t = d.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for n, (k, lab) in enumerate([("humanos", "Humanos"), ("tecnologicos", "Tecnológicos"),
                                  ("materiales", "Materiales")]):
        _celda(t.rows[n].cells[0], lab + ":", True)
        _celda(t.rows[n].cells[1], rec.get(k, ""))
    d.add_paragraph()

    if s.get("perfil_egreso"):
        _h(d, "PERFIL DE EGRESO")
        d.add_paragraph("Al finalizar el curso, el participante será capaz de:")
        _bul(d, s["perfil_egreso"])

    if s.get("alertas"):
        d.add_page_break()
        _h(d, "NOTAS PARA VALIDACIÓN DEL ÁREA ACADÉMICA")
        _bul(d, s["alertas"])

    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Formato Syllabus Cursos 25-26_1.2 · www.clubmentesbrillantes.com.mx")
    r.font.size = Pt(8)
    r.italic = True
    return _guardar(d)


# ---------------------------------------------------------------- otros docs

def marco_docx(m: dict, curso: str) -> bytes:
    d = _doc()
    d.add_heading(curso, 0)
    d.add_paragraph("Marco del curso — apertura y cierre en Moodle")

    for titulo, contenido in (("Avisos", m.get("avisos")),):
        if contenido:
            _h(d, titulo)
            for a in contenido:
                _kv(d, a.get("titulo", ""), a.get("cuerpo", ""))

    b = m.get("bienvenida") or {}
    _h(d, b.get("titulo", "Bienvenida"))
    _p(d, b.get("cuerpo"))

    _h(d, "Requisitos del curso")
    req = m.get("requisitos") or {}
    _p(d, req.get("resumen"))
    for k in ("asistencia", "calificacion_modulo", "acreditacion_programa"):
        if req.get(k):
            _kv(d, k.replace("_", " ").title(), req[k])

    f = m.get("facilitador") or {}
    _h(d, "Conoce a tu facilitador")
    _kv(d, "Impartido por", f.get("nombre", ""))
    if f.get("titular"):
        _kv(d, "Perfil", f["titular"])
    for bl in f.get("bullets", []):
        _kv(d, bl.get("etiqueta", ""), bl.get("contenido", ""))
    if f.get("frase"):
        _kv(d, "Frase", f["frase"], italic=True)

    it = m.get("importancia_tema") or {}
    _h(d, "Importancia del tema")
    _p(d, it.get("pregunta_provocadora"))
    _h(d, "Contras de no aprender el tema", 3)
    _bul(d, it.get("contras_de_no_aprender"))
    _h(d, "Pros de aprender el tema", 3)
    _bul(d, it.get("pros_de_aprender"))
    _kv(d, "Takeaway", it.get("takeaway", ""), italic=True)

    intro = m.get("introduccion") or {}
    _h(d, "Introducción")
    _kv(d, "Objetivo del curso", intro.get("objetivo_curso", ""))
    for c in intro.get("contenido", []):
        _kv(d, f"Módulo {c.get('modulo')}. {c.get('nombre')}", c.get("objetivo", ""))
    if intro.get("audiencia"):
        _h(d, "Este curso está diseñado para", 3)
        _bul(d, intro["audiencia"])

    _h(d, "Recursos del curso")
    for r in m.get("recursos_curso", []):
        _kv(d, f"{r.get('nombre')} ({r.get('tipo')})", r.get("descripcion", ""))

    _h(d, "Material para el instructor")
    for r in m.get("material_instructor", []):
        _kv(d, r.get("nombre", ""), r.get("descripcion", ""))

    _h(d, "Bibliografía y fuentes")
    _bul(d, m.get("bibliografia"))

    ed = m.get("evaluacion_diagnostica_placeholder") or {}
    _h(d, "Evaluación diagnóstica")
    d.add_paragraph(str(ed.get("instrucciones", "")))
    d.add_paragraph("[ESPACIO RESERVADO — ver documento de evaluaciones]")

    ins = m.get("insignia") or {}
    _h(d, "Insignia")
    _p(d, ins.get("resumen_acreditacion"))
    _kv(d, "Dónde consultarla", ins.get("donde_consultar", ""))
    _kv(d, "Qué puedes hacer con ella", ins.get("que_puede_hacer", ""))
    if ins.get("vigencia"):
        _kv(d, "Vigencia", ins["vigencia"])

    cc = m.get("cierre_curso") or {}
    _h(d, "Cierre del curso")
    _p(d, cc.get("texto"))
    _kv(d, "Llamada a la acción", cc.get("llamada_accion", ""))
    _kv(d, "Takeaway", cc.get("takeaway", ""), italic=True)

    enc = m.get("encuesta_satisfaccion") or {}
    _h(d, "Evaluación final del curso y del instructor")
    _kv(d, "Escala", enc.get("escala", ""))
    _h(d, "Sobre el curso", 3)
    _bul(d, enc.get("preguntas_curso"))
    _h(d, "Sobre el instructor", 3)
    _bul(d, enc.get("preguntas_instructor"))

    if m.get("alertas"):
        d.add_page_break()
        _h(d, "Notas para validación")
        _bul(d, m["alertas"])
    return _guardar(d)


def evaluaciones_docx(ev: dict, curso: str) -> bytes:
    d = _doc()
    d.add_heading(curso, 0)
    d.add_paragraph("Sistema de evaluación")

    def banco(titulo, blk):
        _h(d, titulo)
        if blk.get("instrucciones"):
            _p(d, blk["instrucciones"])
        for p in blk.get("preguntas", []):
            _h(d, f"{p.get('id')}. {p.get('enunciado')}", 3)
            if p.get("caso"):
                par = d.add_paragraph()
                r = par.add_run(str(p["caso"]))
                r.italic = True
            for o in p.get("opciones", []):
                marca = "✔ " if o.get("letra") == p.get("respuesta_correcta") else "   "
                par = d.add_paragraph(style="List Bullet")
                r = par.add_run(f"{marca}{o.get('letra')}) {o.get('texto')}")
                if o.get("letra") == p.get("respuesta_correcta"):
                    r.bold = True
            _kv(d, "Explicación", p.get("explicacion", ""))
            _kv(d, "Nivel", p.get("nivel_cognitivo", ""))

    banco("Evaluación diagnóstica", ev.get("diagnostica") or {})
    for b in ev.get("por_modulo", []):
        d.add_page_break()
        banco(f"Evaluación Módulo {b.get('modulo')} — {b.get('titulo', '')}", b)

    gifts = [(("Diagnóstica"), (ev.get("diagnostica") or {}).get("gift"))]
    gifts += [(f"Módulo {b.get('modulo')}", b.get("gift")) for b in ev.get("por_modulo", [])]
    if any(g for _, g in gifts):
        d.add_page_break()
        _h(d, "Bancos en formato GIFT (para importar a Moodle)")
        for nombre, g in gifts:
            if g:
                _h(d, nombre, 3)
                par = d.add_paragraph()
                r = par.add_run(str(g))
                r.font.name = "Consolas"
                r.font.size = Pt(8)
    return _guardar(d)


def presentacion_docx(pres: dict, curso: str) -> bytes:
    d = _doc(0.6)
    d.styles["Normal"].font.size = Pt(9)
    d.add_heading(curso, 0)
    d.add_paragraph(f"Presentación — modalidad {pres.get('modalidad', '')}")

    cols = pres.get("columnas_tabla") or ["Slide", "Título", "Contenido",
                                          "Imagen/diagrama", "Notas del Instructor", "Script"]
    t = d.add_table(rows=1, cols=len(cols))
    t.style = "Table Grid"
    for n, c in enumerate(cols):
        _celda(t.rows[0].cells[n], c, True, 8.5)

    lleva_dur = "Duración" in cols
    for s in pres.get("slides", []):
        row = t.add_row()
        cont = s.get("contenido") or {}
        piezas = []
        if cont.get("bullets"):
            piezas += [f"• {b}" for b in cont["bullets"]]
        for k, etiqueta in (("pros", "Pros"), ("contras", "Contras"), ("preguntas", "Preguntas")):
            if cont.get(k):
                piezas.append(f"{etiqueta}: " + " / ".join(str(x) for x in cont[k]))
        if cont.get("takeaway"):
            piezas.append(f"Takeaway: {cont['takeaway']}")

        valores = [str(s.get("numero", "")), str(s.get("titulo", ""))]
        if lleva_dur:
            valores.append(str(s.get("duracion_min", "") or ""))
        valores += ["\n".join(piezas), str(s.get("imagen_o_diagrama", "")),
                    str(s.get("notas_instructor", "")), str(s.get("script", ""))]
        for n, v in enumerate(valores[:len(cols)]):
            _celda(row.cells[n], v, size=8)

    if pres.get("minuto_a_minuto"):
        d.add_page_break()
        _h(d, "Plan de sesión — minuto a minuto")
        cols2 = ["Hora", "Tema y objetivos", "Tiempo", "Actividad", "Slides"]
        t2 = d.add_table(rows=1, cols=len(cols2))
        t2.style = "Table Grid"
        for n, c in enumerate(cols2):
            _celda(t2.rows[0].cells[n], c, True, 8.5)
        for m in pres["minuto_a_minuto"]:
            row = t2.add_row()
            for n, v in enumerate([m.get("hora", ""), m.get("tema_y_objetivos", ""),
                                   f"{m.get('tiempo_min', '')} min", m.get("actividad", ""),
                                   m.get("slides", "")]):
                _celda(row.cells[n], v, size=8)
    return _guardar(d)


def video_docx(v: dict, curso: str) -> bytes:
    d = _doc()
    d.add_heading(curso, 0)
    d.add_paragraph(f"Material para video — Módulo {v.get('modulo')} "
                    f"({v.get('duracion_total_min', 15)} min)")
    for s in v.get("slides", []):
        _h(d, f"Slide {s.get('numero')} — {s.get('titulo')} "
              f"({s.get('duracion_seg', '')} s)", 2)
        _bul(d, s.get("bullets"))
        _kv(d, "Imagen o diagrama", s.get("imagen_o_diagrama", ""))
        _h(d, "Script", 3)
        _p(d, s.get("script"))
    if v.get("takeaway_modulo"):
        _kv(d, "Takeaway del módulo", v["takeaway_modulo"], italic=True)
    return _guardar(d)


# ------------------------------------------------------------------- paquete

def paquete_zip(resultado: dict) -> bytes:
    """Empaqueta todos los documentos y los JSON en un solo .zip."""
    ficha = resultado.get("ficha") or {}
    syl = resultado.get("syllabus") or {}
    ident = syl.get("identificacion") or {}
    curso = ident.get("curso") or ficha.get("nombre") or "Curso"
    codigo = ident.get("codigo") or ficha.get("codigo") or "SIN-CODIGO"
    base = f"{codigo}".replace("/", "-")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        if syl:
            z.writestr(f"{base}_Syllabus.docx", syllabus_docx(syl))
        if resultado.get("marco"):
            z.writestr(f"{base}_Marco_del_curso.docx",
                       marco_docx(resultado["marco"], curso))
        for k, v in sorted(resultado.items()):
            if k.startswith("modulo_") and isinstance(v, dict):
                z.writestr(f"{base}_Modulo_{v.get('numero', k.split('_')[-1])}.docx",
                           modulo_docx(v, curso, codigo))
            if k.startswith("video_") and isinstance(v, dict):
                z.writestr(f"{base}_Video_Modulo_{v.get('modulo', '')}.docx",
                           video_docx(v, curso))
            if k.startswith("presentacion_") and isinstance(v, dict):
                z.writestr(f"{base}_Presentacion_{k.split('_', 1)[1]}.docx",
                           presentacion_docx(v, curso))
        if resultado.get("evaluaciones"):
            z.writestr(f"{base}_Evaluaciones.docx",
                       evaluaciones_docx(resultado["evaluaciones"], curso))

        for k, v in resultado.items():
            if k in ("hallazgos", "alertas", "uso", "errores") or isinstance(v, dict):
                z.writestr(f"json/{k}.json",
                           json.dumps(v, ensure_ascii=False, indent=2))

        z.writestr("_generado.txt",
                   f"Curso: {curso}\nCódigo: {codigo}\n"
                   f"Generado: {datetime.now():%Y-%m-%d %H:%M}\n"
                   f"Generador de Cursos — Club de Mentes Brillantes\n")
    return buf.getvalue()
