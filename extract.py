"""Extracción de texto de la ficha técnica: .docx, .pdf, .txt o .md."""

import io
import re
from pathlib import Path


def _from_docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    partes = []

    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        estilo = p.style.name or ""
        if estilo.startswith("Heading") or estilo.startswith("Título"):
            partes.append(f"\n## {t}")
        else:
            partes.append(t)

    for i, tabla in enumerate(d.tables):
        partes.append(f"\n[TABLA {i + 1}]")
        for fila in tabla.rows:
            celdas = [c.text.strip().replace("\n", " ") for c in fila.cells]
            # las tablas de Word repiten la celda en merges; se colapsa
            limpias, previa = [], None
            for c in celdas:
                if c != previa:
                    limpias.append(c)
                previa = c
            if any(limpias):
                partes.append(" | ".join(limpias))

    return "\n".join(partes)


def _from_pdf(data: bytes) -> str:
    import pdfplumber
    partes = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for n, page in enumerate(pdf.pages, 1):
            texto = page.extract_text() or ""
            if texto.strip():
                partes.append(texto)
            for tabla in page.extract_tables() or []:
                partes.append(f"\n[TABLA pág. {n}]")
                for fila in tabla:
                    celdas = [(c or "").strip().replace("\n", " ") for c in fila]
                    if any(celdas):
                        partes.append(" | ".join(celdas))
    return "\n".join(partes)


def extraer(nombre: str, data: bytes) -> str:
    """Devuelve el texto plano de la ficha, sea cual sea su formato."""
    ext = Path(nombre).suffix.lower()

    if ext == ".docx":
        texto = _from_docx(data)
    elif ext == ".pdf":
        texto = _from_pdf(data)
    elif ext in (".txt", ".md"):
        texto = data.decode("utf-8", errors="replace")
    else:
        raise ValueError(
            f"Formato no soportado: {ext or 'sin extensión'}. "
            "Se aceptan .docx, .pdf, .txt y .md."
        )

    texto = re.sub(r"\n{3,}", "\n\n", texto).strip()

    if len(texto) < 200:
        raise ValueError(
            "El documento tiene muy poco texto extraíble "
            f"({len(texto)} caracteres). Si es un PDF escaneado, necesita OCR previo."
        )
    return texto


def resumen_extraccion(texto: str) -> dict:
    """Métricas rápidas para mostrar en la UI antes de generar."""
    palabras = len(re.findall(r"\b\w+\b", texto))
    return {
        "caracteres": len(texto),
        "palabras": palabras,
        "lineas": texto.count("\n") + 1,
        "tiene_tablas": "[TABLA" in texto,
        "preview": texto[:1200],
    }
